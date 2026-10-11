"""Text extraction with page-level provenance.

Page numbers are preserved because they are what a citation points at
("[Source: 10-K, page 42]"). Formats without pages get synthetic page boundaries so every
chunk still has a stable, human-checkable locator.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Any

from finagent_core.errors import DocumentParseError, UnsupportedFileTypeError
from finagent_core.logging import get_logger

logger = get_logger(__name__)

#: Characters per synthetic page for formats with no intrinsic pagination.
SYNTHETIC_PAGE_CHARS = 3000


@dataclass
class ExtractedPage:
    page_number: int
    text: str


@dataclass
class ExtractedDocument:
    pages: list[ExtractedPage]
    metadata: dict[str, Any] = field(default_factory=dict)
    tabular_rows: list[dict[str, Any]] = field(default_factory=list)
    pagination: str = "native"

    @property
    def page_count(self) -> int:
        return len(self.pages)

    @property
    def text(self) -> str:
        return "\n\n".join(page.text for page in self.pages)

    @property
    def char_count(self) -> int:
        return sum(len(page.text) for page in self.pages)


def _paginate(text: str, *, chars_per_page: int = SYNTHETIC_PAGE_CHARS) -> list[ExtractedPage]:
    """Split plain text into synthetic pages at paragraph boundaries."""
    paragraphs = [block.strip() for block in text.split("\n\n") if block.strip()]
    pages: list[ExtractedPage] = []
    buffer: list[str] = []
    length = 0
    for paragraph in paragraphs:
        if length + len(paragraph) > chars_per_page and buffer:
            pages.append(ExtractedPage(len(pages) + 1, "\n\n".join(buffer)))
            buffer, length = [], 0
        buffer.append(paragraph)
        length += len(paragraph)
    if buffer:
        pages.append(ExtractedPage(len(pages) + 1, "\n\n".join(buffer)))
    return pages or [ExtractedPage(1, text.strip())]


def extract_pdf(payload: bytes) -> ExtractedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise DocumentParseError("pypdf is not installed") from exc
    try:
        reader = PdfReader(io.BytesIO(payload))
        pages = [
            ExtractedPage(index + 1, (page.extract_text() or "").strip())
            for index, page in enumerate(reader.pages)
        ]
    except Exception as exc:  # pypdf raises a wide variety of errors
        raise DocumentParseError(f"pdf parse failure: {exc}") from exc

    populated = [page for page in pages if page.text]
    if not populated:
        raise DocumentParseError(
            "no extractable text found in PDF (scanned documents require OCR, which is not configured)"
        )
    return ExtractedDocument(
        pages=pages,
        metadata={
            "page_count": len(pages),
            "extractor": "pypdf",
            "empty_pages": len(pages) - len(populated),
        },
    )


def extract_docx(payload: bytes) -> ExtractedDocument:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("python-docx is not installed") from exc
    try:
        document = docx.Document(io.BytesIO(payload))
    except Exception as exc:
        raise DocumentParseError(f"docx parse failure: {exc}") from exc

    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    if not blocks:
        raise DocumentParseError("docx contains no text")
    return ExtractedDocument(
        pages=_paginate("\n\n".join(blocks)),
        metadata={"extractor": "python-docx", "paragraphs": len(blocks)},
        pagination="synthetic",
    )


def extract_text(payload: bytes) -> ExtractedDocument:
    try:
        text = payload.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = payload.decode("latin-1")
        except UnicodeDecodeError as exc:
            raise DocumentParseError("could not decode text file") from exc
    if not text.strip():
        raise DocumentParseError("text file is empty")
    return ExtractedDocument(pages=_paginate(text), metadata={"extractor": "text"}, pagination="synthetic")


def extract_csv(payload: bytes) -> ExtractedDocument:
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = payload.decode("latin-1")
    try:
        sample = text[:8192]
        dialect = csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
    except csv.Error:
        dialect = csv.excel
    rows = [dict(row) for row in csv.DictReader(io.StringIO(text), dialect=dialect)]
    if not rows:
        raise DocumentParseError("csv contains no data rows")

    headers = list(rows[0].keys())
    lines = [" | ".join(str(header) for header in headers)]
    lines += [" | ".join(str(row.get(header, "")) for header in headers) for row in rows]
    return ExtractedDocument(
        pages=_paginate("\n\n".join(lines)),
        metadata={"extractor": "csv", "row_count": len(rows), "columns": headers},
        tabular_rows=rows,
        pagination="synthetic",
    )


def extract_xlsx(payload: bytes) -> ExtractedDocument:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("openpyxl is not installed") from exc
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentParseError(f"xlsx parse failure: {exc}") from exc

    pages: list[ExtractedPage] = []
    rows_out: list[dict[str, Any]] = []
    for index, sheet in enumerate(workbook.worksheets, start=1):
        values = [
            [("" if cell is None else str(cell)) for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        if not values:
            continue
        header = values[0]
        for row in values[1:]:
            if any(cell != "" for cell in row):
                rows_out.append(
                    {header[i] if i < len(header) else f"col_{i}": row[i] for i in range(len(row))}
                )
        pages.append(
            ExtractedPage(index, "\n".join([f"Sheet: {sheet.title}"] + [" | ".join(r) for r in values]))
        )
    workbook.close()
    if not pages:
        raise DocumentParseError("workbook contains no data")
    return ExtractedDocument(
        pages=pages,
        metadata={"extractor": "openpyxl", "sheets": len(pages), "row_count": len(rows_out)},
        tabular_rows=rows_out,
    )


_EXTRACTORS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "txt": extract_text,
    "md": extract_text,
    "csv": extract_csv,
    "xlsx": extract_xlsx,
}


def extract_document(payload: bytes, extension: str) -> ExtractedDocument:
    extractor = _EXTRACTORS.get(extension.lower())
    if extractor is None:
        raise UnsupportedFileTypeError(f"no extractor for .{extension}")
    document = extractor(payload)
    logger.info(
        "document_extracted",
        extra={
            "extension": extension,
            "page_count": document.page_count,
            "char_count": document.char_count,
        },
    )
    return document
